"""Generate the PBL platform submission document (Word .docx).

Run:  python scripts/make-submission.py
Out:  EthiCross-Submission.docx  (repo root)

Fill in GROUP_NUMBER below before running, and confirm PLATFORM_LINK
after enabling GitHub Pages (Settings -> Pages -> Source: GitHub Actions).
"""
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

GROUP_NUMBER = "___"  # <-- TODO: fill in your group number, e.g. "5"
PLATFORM_LINK = "https://differXI.github.io/EthiCross/"

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

# ---------- Title block ----------
t = doc.add_heading("EthiCross — PBL Platform Submission", level=0)
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run(
    "Ethics and Professionalism for Software Engineers (953420)\n"
    "College of Arts, Media, and Technology, Chiang Mai University\n"
    "Instructors: Asst. Prof. Dr. Pradorn Sureephong, "
    "Asst. Prof. Dr. Suepphong Chernbumroong"
)
r.font.size = Pt(11)


def h1(text):
    doc.add_heading(text, level=1)


def h2(text):
    doc.add_heading(text, level=2)


def para(text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
    p.add_run(text)
    return p


def bullets(items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


# ---------- 1-3: identity ----------
h1("1. Group Number")
para(f"Group {GROUP_NUMBER}")

h1("2. Project / Platform Name")
para("EthiCross — an educational crossword game for learning software-ethics terminology.")

h1("3. Platform Link")
p = doc.add_paragraph()
r = p.add_run(PLATFORM_LINK)
r.bold = True
para(
    "If the link does not open yet, enable hosting once: GitHub repository "
    "Settings → Pages → Build and deployment → Source: “GitHub Actions”. "
    "Every push to the main branch then redeploys the platform automatically."
)

# ---------- 4: access ----------
h1("4. Access / Login Instructions")
bullets(
    [
        "No account, no login, and no special permission are required.",
        "Open the Platform Link above on a mobile phone, tablet, or computer with internet access.",
        "All modes needed for user testing are available immediately: Levels, Endless, Daily, "
        "Study Index, Knowledge Test, and theme selection.",
        "Progress (levels cleared, stars, streaks, test results) is stored in the browser on the user's own device.",
    ]
)

# ---------- 5: description ----------
h1("5. Brief Description of the Platform")
para(
    "EthiCross turns software-ethics terminology into interactive play. Instead of memorizing "
    "definitions from slides, students solve crossword clues about real software-engineering "
    "situations. Every correct answer unlocks a short lesson (definition, why it matters, and a "
    "real-world example), following the loop Learn → Play → Feedback → Understand → Review."
)
para(
    "The platform teaches 70 ethics terms — privacy, security, fairness, accountability, "
    "transparency, copyright, accessibility, plus data-care, AI-fairness, professional-duty, and "
    "open-source topics such as consent, encryption, bias, plagiarism, and liability."
)
h2("5.1 Main functions (all operational for user testing)")
rows = [
    ("🏁 Levels", "100 crosswords in 20 themed chapters (4→7 words, warm-up→expert) with locks, stars, and score/time tracking."),
    ("♾️ Endless", "Survival mode: clues keep coming until 3 wrong answers cost all hearts; streaks and best score."),
    ("📅 Daily", "One fresh crossword per day for everyone, one play per day, with day-streaks."),
    ("📚 Study Index", "Read-before-you-play glossary: all 70 terms grouped by world, with search and ✓ learned marks."),
    ("📝 Knowledge Test", "Required post-test and optional pre-test (10 questions each), history, and Copy-CSV export for data collection."),
    ("🎨 Themes", "Four visual themes (Scholar default, Candy, Neon, Pixel), switchable at any time."),
]
table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text = "Function"
hdr[1].text = "What the user can do"
for name, desc in rows:
    cells = table.add_row().cells
    cells[0].text = name
    cells[1].text = desc
para(
    "Scoring is identical everywhere: +100 per correct answer, −25 per hint, "
    "+500 completion bonus. Answers are case-insensitive and a full answer is never revealed after a wrong attempt."
)

# ---------- 6: objectives ----------
h1("6. Learning Objectives")
para("After using EthiCross, students will be able to:")
doc.add_paragraph("Define core software-ethics terms (privacy, security, fairness, accountability, and more).", style="List Number")
doc.add_paragraph("Explain why each principle matters in real software-engineering work.", style="List Number")
doc.add_paragraph("Recognize ethical issues in realistic scenarios, from dark patterns to data breaches.", style="List Number")
doc.add_paragraph("Apply ethical reasoning to judge design decisions — and demonstrate it in the Knowledge Test.", style="List Number")

# ---------- 7: knowledge test ----------
h1("7. Knowledge Test (Post-test required; Pre-test optional)")
h2("7.1 Format (built into the platform)")
bullets(
    [
        "Location: Home → 📝 Knowledge Test (no login needed).",
        "Each test has 10 multiple-choice questions drawn from the 70-term bank: the player reads a definition and picks the correct term from 4 options.",
        "Post-test (required): taken after playing — measures what was learned.",
        "Pre-test (optional): taken before playing — enables before/after comparison.",
        "After submitting, the player sees score, percentage, and a full per-question review.",
    ]
)
h2("7.2 Data collection and reporting")
bullets(
    [
        "Every attempt is stored on the device with test type, optional name/ID, score, and date.",
        "The “📋 Copy CSV” button exports all attempts as test,name,score,total,percent,date rows for the group's spreadsheet and course report.",
        "Suggested testing flow for each participant: (1) optional Pre-test → (2) play Levels and/or read the Study Index → (3) required Post-test → (4) Copy CSV and send it to the group.",
        "Per the course agenda, results are reported as actually observed; no particular improvement is required.",
    ]
)

# ---------- POC mapping ----------
h1("8. How the Platform Meets the POC Requirements")
reqs = [
    ("§2.1 Learning objectives", "Objectives are shown in the app (§6 above) and each game action maps to one: solving (Define), lessons (Explain), examples (Recognize), tests (Apply)."),
    ("§2.2 Accessibility", "Public link, no login, works on phones and computers."),
    ("§2.3 Bug fixing", "Core loop tested (lint + production build pass); the group continues fixing issues from tester feedback without changing core objectives."),
    ("§2.4 Knowledge test", "Post-test required and pre-test optional, both in-app with exportable results (§7 above)."),
]
table2 = doc.add_table(rows=1, cols=2)
table2.style = "Table Grid"
hdr2 = table2.rows[0].cells
hdr2[0].text = "Agenda requirement"
hdr2[1].text = "How EthiCross satisfies it"
for req, how in reqs:
    cells = table2.add_row().cells
    cells[0].text = req
    cells[1].text = how

# ---------- team ----------
h1("9. Team Responsibilities")
members = [
    ("Wanikkasit Nopthiraitthikun", "UI/UX Designer & Frontend Developer"),
    ("Chonchanun Khachonphurithanakul", "System Tester & Quality Assurance"),
    ("Thanachai Naksomboon", "Backend Developer & Database Administrator"),
    ("Panuwat Songkram", "Frontend Developer & Game Logic"),
    ("Ratthasas Singhamanee", "Project Manager & Documentation"),
]
table3 = doc.add_table(rows=1, cols=2)
table3.style = "Table Grid"
hdr3 = table3.rows[0].cells
hdr3[0].text = "Member"
hdr3[1].text = "Responsibility"
for name, role in members:
    cells = table3.add_row().cells
    cells[0].text = name
    cells[1].text = role

doc.save("EthiCross-Submission.docx")
print("Saved EthiCross-Submission.docx")
